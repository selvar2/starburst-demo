# Additional dependencies: fastapi, uvicorn[standard], openpyxl, python-multipart

import io
import os
import csv
import re
import time
import threading
from concurrent.futures import ThreadPoolExecutor, as_completed
from numbers import Number
from pathlib import Path

from fastapi import FastAPI, HTTPException, Response
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

from starburst_client_jwt import StarburstClientJWT as StarburstClient
from permission_manager import PermissionManager
from analytics_service import build_result_insights
from llm_provider import build_provider_from_env
from nl2sql_service import NL2SQLService

# ---------------------------------------------------------------------------
# App & client
# ---------------------------------------------------------------------------

app = FastAPI(title="Starburst Galaxy Chatbot (JWT)")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

client = StarburstClient()
_perms = PermissionManager()
_developer = os.getenv("STARBURST_DEVELOPER", "unknown")
_llm_provider = build_provider_from_env()

HERE = Path(__file__).parent

# Static files (best-effort; directory may not exist yet)
static_dir = HERE / "static"
if static_dir.is_dir():
    app.mount("/static", StaticFiles(directory=str(static_dir)), name="static")


# ---------------------------------------------------------------------------
# Models
# ---------------------------------------------------------------------------

class QueryRequest(BaseModel):
    sql: str

class ChatRequest(BaseModel):
    message: str
    context: dict | None = None

class ExportRequest(BaseModel):
    columns: list[str]
    rows: list[list]
    title: str = "export"


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _catalog(ctx: dict | None) -> str:
    return (ctx or {}).get("catalog") or client.catalog or "mcp2ohio"

def _schema(ctx: dict | None) -> str:
    return (ctx or {}).get("schema") or client.schema or "test_writes"

class QueryError(Exception):
    pass

_AGG_ALIAS_RE = re.compile(
    r'\b(COUNT|SUM|AVG|MIN|MAX)\s*\(\s*(\*|\w+)\s*\)(?!\s+AS\b)',
    re.IGNORECASE,
)

def _add_agg_aliases(sql: str) -> str:
    """Add aliases to aggregate functions that don't have one."""
    def _repl(m):
        func = m.group(1).lower()
        col = m.group(2)
        if col == '*':
            alias = f"{func}"
        else:
            alias = f"{func}_{col}"
        return f'{m.group(0)} AS {alias}'
    return _AGG_ALIAS_RE.sub(_repl, sql)

def _exec(sql: str) -> dict:
    """Execute and normalise result."""
    try:
        result = client.execute(_add_agg_aliases(sql))
    except Exception as exc:
        raise QueryError(str(exc))
    if "rows_affected" in result:
        return {
            "columns": ["rows_affected"],
            "rows": [[result["rows_affected"]]],
            "row_count": 1,
            "rows_affected": result["rows_affected"],
            "status": result.get("status", "ok"),
        }
    columns = result.get("columns", [])
    rows = result.get("rows", [])
    return {"columns": columns, "rows": rows, "row_count": len(rows)}


def _verification_target_fqn(target: dict) -> str:
    return f"{target['catalog']}.{target['schema']}.{target['table']}"


def _split_sql_list(text: str) -> list[str]:
    parts: list[str] = []
    buf: list[str] = []
    depth = 0
    in_quote = False
    idx = 0
    while idx < len(text):
        ch = text[idx]
        if ch == "'":
            buf.append(ch)
            if in_quote and idx + 1 < len(text) and text[idx + 1] == "'":
                buf.append(text[idx + 1])
                idx += 2
                continue
            in_quote = not in_quote
            idx += 1
            continue
        if not in_quote:
            if ch == "(":
                depth += 1
            elif ch == ")" and depth > 0:
                depth -= 1
            elif ch == "," and depth == 0:
                part = "".join(buf).strip()
                if part:
                    parts.append(part)
                buf = []
                idx += 1
                continue
        buf.append(ch)
        idx += 1
    tail = "".join(buf).strip()
    if tail:
        parts.append(tail)
    return parts


def _extract_parenthesized_groups(text: str) -> list[str]:
    groups: list[str] = []
    depth = 0
    in_quote = False
    start: int | None = None
    idx = 0
    while idx < len(text):
        ch = text[idx]
        if ch == "'":
            if in_quote and idx + 1 < len(text) and text[idx + 1] == "'":
                idx += 2
                continue
            in_quote = not in_quote
            idx += 1
            continue
        if not in_quote:
            if ch == "(":
                if depth == 0:
                    start = idx + 1
                depth += 1
            elif ch == ")" and depth > 0:
                depth -= 1
                if depth == 0 and start is not None:
                    groups.append(text[start:idx].strip())
                    start = None
        idx += 1
    return groups


def _parse_identifier_list(text: str) -> list[str]:
    return [part.strip().strip('"') for part in _split_sql_list(text)]


def _parse_assignment_map(text: str) -> dict[str, str] | None:
    assignments: dict[str, str] = {}
    for part in _split_sql_list(text):
        m = re.match(r'\s*"?(?P<col>[a-zA-Z_]\w*)"?\s*=\s*(?P<value>.+?)\s*$', part, re.IGNORECASE | re.DOTALL)
        if not m:
            return None
        assignments[m.group("col")] = m.group("value").strip()
    return assignments or None


def _parse_where_equalities(text: str) -> dict[str, str] | None:
    predicates: dict[str, str] = {}
    for part in re.split(r'\bAND\b', text, flags=re.IGNORECASE):
        clause = part.strip().strip("() ")
        m = re.match(r'\s*"?(?P<col>[a-zA-Z_]\w*)"?\s*=\s*(?P<value>.+?)\s*$', clause, re.IGNORECASE | re.DOTALL)
        if not m:
            return None
        predicates[m.group("col")] = m.group("value").strip()
    return predicates or None


def _describe_table_columns(target: dict) -> list[str]:
    fqn = _verification_target_fqn(target)
    result = client.execute(f"DESCRIBE {fqn}")
    columns: list[str] = []
    for row in result.get("rows", []):
        if row and row[0]:
            name = str(row[0]).strip().strip('"')
            if name and not name.startswith("#"):
                columns.append(name)
    return columns


def _build_insert_verification_sql(sql: str, target: dict) -> tuple[str | None, str | None]:
    match = re.search(r'\bINSERT\s+INTO\s+[\w."]+\s*(?:\((?P<columns>.*?)\))?\s*VALUES\s*(?P<values>.+?)\s*;?\s*$', sql, re.IGNORECASE | re.DOTALL)
    if not match:
        return None, "Could not parse INSERT statement for verification."
    columns = _parse_identifier_list(match.group("columns")) if match.group("columns") else _describe_table_columns(target)
    row_groups = _extract_parenthesized_groups(match.group("values"))
    if not columns:
        return None, "Unable to infer inserted columns for verification."
    if not row_groups:
        return None, "Could not parse inserted VALUES for verification."
    predicates: list[str] = []
    for group in row_groups:
        values = _split_sql_list(group)
        if len(values) != len(columns):
            return None, "Inserted values do not match the target columns."
        predicates.append("(" + " AND ".join(f"{col} = {value.strip()}" for col, value in zip(columns, values)) + ")")
    return f"SELECT * FROM {_verification_target_fqn(target)} WHERE " + " OR ".join(predicates) + " LIMIT 100", None


def _build_update_verification_sql(sql: str, target: dict) -> tuple[str | None, str | None]:
    match = re.search(r'\bUPDATE\s+[\w."]+\s+SET\s+(?P<set>.+?)\s+WHERE\s+(?P<where>.+?)\s*;?\s*$', sql, re.IGNORECASE | re.DOTALL)
    if not match:
        return None, "Could not parse UPDATE statement for verification."
    set_values = _parse_assignment_map(match.group("set"))
    where_values = _parse_where_equalities(match.group("where"))
    if not set_values:
        return None, "Could not parse updated columns for verification."
    if not where_values:
        return None, "Could not safely parse the UPDATE WHERE clause for verification."
    verification_values = dict(where_values)
    verification_values.update(set_values)
    predicate_sql = " AND ".join(f"{column} = {value}" for column, value in verification_values.items())
    return f"SELECT * FROM {_verification_target_fqn(target)} WHERE {predicate_sql} LIMIT 100", None


def _sql_string_literal(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


def _build_create_table_verification_sql(target: dict) -> tuple[str | None, str | None]:
    if not target.get("table"):
        return None, "CREATE TABLE verification requires a fully qualified table target."
    return (
        f"SELECT table_catalog, table_schema, table_name FROM {target['catalog']}.information_schema.tables "
        f"WHERE table_schema = {_sql_string_literal(target['schema'])} "
        f"AND table_name = {_sql_string_literal(target['table'])} LIMIT 100",
        None,
    )


def _build_create_schema_verification_sql(target: dict) -> tuple[str | None, str | None]:
    return (
        f"SELECT catalog_name, schema_name FROM {target['catalog']}.information_schema.schemata "
        f"WHERE schema_name = {_sql_string_literal(target['schema'])} LIMIT 100",
        None,
    )


def _build_write_verification_sql(sql: str, op_label: str, target: dict) -> tuple[str | None, str | None]:
    if op_label == "INSERT":
        return _build_insert_verification_sql(sql, target)
    if op_label == "UPDATE":
        return _build_update_verification_sql(sql, target)
    if op_label == "CREATE TABLE":
        return _build_create_table_verification_sql(target)
    if op_label == "CREATE SCHEMA":
        return _build_create_schema_verification_sql(target)
    return None, "Verification queries are only supported for INSERT, UPDATE, CREATE TABLE, and CREATE SCHEMA."


def _attach_write_verification(result: dict, sql: str, op_label: str, target: dict) -> dict:
    if op_label not in {"INSERT", "UPDATE", "CREATE TABLE", "CREATE SCHEMA"}:
        return result
    if op_label in {"INSERT", "UPDATE"} and "rows_affected" not in result:
        return result
    enriched = dict(result)
    verification_sql, verification_reason = _build_write_verification_sql(sql, op_label, target)
    if not verification_sql:
        enriched["verification_reason"] = verification_reason
        return enriched
    enriched["verification_sql"] = verification_sql
    try:
        verification_result = _exec(verification_sql)
    except QueryError as exc:
        enriched["verification_reason"] = str(exc)
        return enriched
    enriched["verification_columns"] = verification_result["columns"]
    enriched["verification_rows"] = verification_result["rows"]
    enriched["verification_row_count"] = verification_result["row_count"]
    enriched["columns"] = verification_result["columns"]
    enriched["rows"] = verification_result["rows"]
    enriched["row_count"] = verification_result["row_count"]
    return enriched


# ---------------------------------------------------------------------------
# DDL/DML classification, FQ validation, permission, destructive guard
# ---------------------------------------------------------------------------

class FQValidationError(Exception):
    pass

class PermissionDenied(Exception):
    pass

class ConfirmRequired(Exception):
    pass


_DESTRUCTIVE_PERMS = {"drop_table", "drop_schema", "truncate", "delete"}

_CLASSIFY_PREFIXES = [
    ("CREATE SCHEMA", "create_schema"),
    ("CREATE TABLE",  "create_table"),
    ("DROP SCHEMA",   "drop_schema"),
    ("DROP TABLE",    "drop_table"),
    ("INSERT",        "insert"),
    ("UPDATE",        "update"),
    ("DELETE",        "delete"),
    ("MERGE",         "merge"),
    ("TRUNCATE",      "truncate"),
    ("ALTER",         "execute_raw"),
    ("SELECT",        "read"),
    ("SHOW",          "read"),
    ("DESCRIBE",      "read"),
    ("WITH",          "read"),
]

def _classify_sql(sql: str) -> tuple[str, str]:
    """Return (op_label, permission_key). op_label is uppercase keyword(s)."""
    s = re.sub(r"\s+", " ", sql.strip().upper())
    for prefix, perm in _CLASSIFY_PREFIXES:
        if s.startswith(prefix):
            return prefix, perm
    return "UNKNOWN", "execute_raw"


_TARGET_PATTERNS = {
    "INSERT":         re.compile(r'\bINSERT\s+INTO\s+([\w."]+)', re.IGNORECASE),
    "UPDATE":         re.compile(r'\bUPDATE\s+([\w."]+)',         re.IGNORECASE),
    "DELETE":         re.compile(r'\bDELETE\s+FROM\s+([\w."]+)',  re.IGNORECASE),
    "TRUNCATE":       re.compile(r'\bTRUNCATE\s+(?:TABLE\s+)?([\w."]+)', re.IGNORECASE),
    "DROP TABLE":     re.compile(r'\bDROP\s+TABLE\s+(?:IF\s+EXISTS\s+)?([\w."]+)', re.IGNORECASE),
    "DROP SCHEMA":    re.compile(r'\bDROP\s+SCHEMA\s+(?:IF\s+EXISTS\s+)?([\w."]+)', re.IGNORECASE),
    "CREATE SCHEMA":  re.compile(r'\bCREATE\s+SCHEMA\s+(?:IF\s+NOT\s+EXISTS\s+)?([\w."]+)', re.IGNORECASE),
    "CREATE TABLE":   re.compile(r'\bCREATE\s+TABLE\s+(?:IF\s+NOT\s+EXISTS\s+)?([\w."]+)', re.IGNORECASE),
    "MERGE":          re.compile(r'\bMERGE\s+INTO\s+([\w."]+)',   re.IGNORECASE),
    "ALTER":          re.compile(r'\bALTER\s+(?:TABLE|SCHEMA)\s+([\w."]+)', re.IGNORECASE),
}

# Schema-level ops require 2-part catalog.schema; everything else 3-part catalog.schema.table.
_SCHEMA_LEVEL_OPS = {"CREATE SCHEMA", "DROP SCHEMA"}
# Read ops are FQ-exempt (SELECT may use unqualified or aliased forms).
_FQ_EXEMPT_OPS = {"SELECT", "SHOW", "DESCRIBE", "WITH", "UNKNOWN"}


def _validate_fq(sql: str, op_label: str) -> dict:
    """Ensure DDL/DML target is fully qualified.

    Returns {catalog, schema, table?}. Raises FQValidationError on failure.
    Read ops return {} (unchecked).
    """
    if op_label in _FQ_EXEMPT_OPS:
        return {}
    pat = _TARGET_PATTERNS.get(op_label)
    if pat is None:
        raise FQValidationError(
            f"Unsupported operation '{op_label}' — refusing to execute without an explicit target."
        )
    m = pat.search(sql)
    if not m:
        raise FQValidationError(
            f"Could not locate target object in {op_label} statement."
        )
    target = m.group(1)
    parts = [p.strip('"') for p in target.split(".")]
    expected = 2 if op_label in _SCHEMA_LEVEL_OPS else 3
    if len(parts) != expected:
        shape = "catalog.schema" if expected == 2 else "catalog.schema.table"
        raise FQValidationError(
            f"{op_label} requires fully qualified {shape} (got '{target}'). "
            f"Specify all parts, e.g. {op_label} mcp2ohio.test_writes" + ("" if expected == 2 else ".my_table")
        )
    for p in parts:
        try:
            StarburstClient.validate_identifier(p)
        except ValueError as e:
            raise FQValidationError(str(e))
    out = {"catalog": parts[0], "schema": parts[1]}
    if expected == 3:
        out["table"] = parts[2]
    return out


def _check_permission(perm_key: str) -> None:
    allowed, msg = _perms.check_with_message(_developer, perm_key)
    if not allowed:
        raise PermissionDenied(msg)


def _delete_has_where(sql: str) -> bool:
    return bool(re.search(r"\bWHERE\b", sql or "", re.IGNORECASE))


def _enforce_destructive_confirm(perm_key: str, op_label: str, target: dict, confirm: bool, sql: str = "") -> None:
    if perm_key not in _DESTRUCTIVE_PERMS:
        return
    if op_label == "DELETE" and not _delete_has_where(sql):
        raise FQValidationError(
            "DELETE requires a WHERE clause. Full-table delete is blocked by policy even when confirmation is supplied."
        )
    if confirm:
        return
    fq = ".".join(p for p in (target.get("catalog"), target.get("schema"), target.get("table")) if p)
    raise ConfirmRequired(
        f"WARNING: {op_label} on '{fq}' is destructive and cannot be undone. "
        f"Resend with context.confirm=true to proceed."
    )


# ---------------------------------------------------------------------------
# Chart suggestion
# ---------------------------------------------------------------------------

_DATE_KEYWORDS = {"date", "time", "timestamp", "day", "month", "year", "created", "updated"}

def _is_numeric_value(value) -> bool:
    return isinstance(value, Number) and not isinstance(value, bool)


def _final_chart_suggestion(llm_chart: str | None, suggested: str) -> str:
    allowed = {"bar", "pie", "line", "area", "table", "metric"}
    if llm_chart not in allowed:
        return suggested
    if llm_chart in {"table", "metric"} and suggested not in {"table", "metric"}:
        return suggested
    return llm_chart

def _suggest_chart(columns: list[str], rows: list[list]) -> str:
    if not columns or not rows:
        return "table"

    ncols = len(columns)

    # Single scalar value
    if ncols == 1 and len(rows) == 1:
        return "metric"

    # Check for date column
    lower_cols = [c.lower() for c in columns]
    has_date = any(any(kw in c for kw in _DATE_KEYWORDS) for c in lower_cols)
    if has_date and ncols >= 2:
        return "line"

    # 2 columns: any type + number (group by results)
    if ncols == 2 and rows:
        second_num = _is_numeric_value(rows[0][1]) if rows[0] else False
        if second_num:
            return "pie" if len(rows) < 8 else "bar"

    # Aggregation (single row, multiple numeric cols)
    if len(rows) == 1 and ncols >= 1:
        if all(_is_numeric_value(v) for v in rows[0]):
            return "bar"

    # Multiple rows with at least one numeric column
    if ncols >= 2 and len(rows) > 1:
        has_num = any(_is_numeric_value(rows[0][i]) for i in range(ncols))
        if has_num:
            return "bar"

    return "table"


# ---------------------------------------------------------------------------
# NL -> SQL
# ---------------------------------------------------------------------------

_SQL_START = re.compile(
    r"^\s*(SELECT|SHOW|DESCRIBE|WITH|INSERT|UPDATE|DELETE|CREATE|DROP|ALTER)\b",
    re.IGNORECASE,
)

def _extract_context(text: str, default_catalog: str, default_schema: str):
    """Extract catalog, schema, table from NL text with context hints."""
    low = text.lower()
    cat = default_catalog
    sch = default_schema

    # Match: "X is part of Y schema and part of Z catalog"
    m = re.search(r'part of\s+(\w+)\s+schema', low)
    if m:
        sch = m.group(1)
    m = re.search(r'part of\s+(\w+)\s+catalog', low)
    if m:
        cat = m.group(1)

    # Match: "from/in Y schema in/of Z catalog"
    m = re.search(r'(?:from|in)\s+(\w+)\s+schema\s+(?:in|of)\s+(\w+)\s+catalog', low)
    if m:
        sch = m.group(1)
        cat = m.group(2)

    # Match: "in Y schema" alone
    m = re.search(r'in\s+(\w+)\s+schema', low)
    if m:
        sch = m.group(1)

    # Match: "in Z catalog" alone
    m = re.search(r'in\s+(\w+)\s+catalog', low)
    if m:
        cat = m.group(1)

    return cat, sch


def _clean_nl(text: str) -> str:
    """Remove context hints from NL text to get core query."""
    cleaned = re.sub(r',?\s*\w+\s+is\s+part\s+of\s+\w+\s+schema\s+and\s+part\s+of\s+\w+\s+catalog\s*', '', text, flags=re.IGNORECASE)
    cleaned = re.sub(r',?\s*(?:from|in)\s+\w+\s+schema\s+(?:in|of)\s+\w+\s+catalog\s*', '', cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r',?\s*in\s+\w+\s+(?:schema|catalog)\s*', '', cleaned, flags=re.IGNORECASE)
    return cleaned.strip()


# ---------------------------------------------------------------------------
# Business-user NL INSERT
#
# Recognizes phrasings like:
#   "Add a new row into catalog X, schema Y, table Z with col1 v1, col2 v2"
#   "Insert a record into X.Y.Z where col1 = v1, col2 = v2, and col3 = v3"
#   "In catalog X, schema Y, add a row to table Z with col1 v1, col2 v2"
# Requires explicit catalog + schema + table — no default fallback (per spec:
# "do not guess missing database object context").
# ---------------------------------------------------------------------------

_NL_INSERT_TRIGGER = re.compile(
    r"^\s*(?:"
    r"add\s+(?:a\s+|the\s+)?(?:new\s+)?(?:row|record)|"
    r"insert\s+(?:a\s+|the\s+)?(?:new\s+)?(?:row|record)|"
    r"in\s+catalog\s+\w+.*?\b(?:add|insert)\s+(?:a\s+|the\s+)?(?:new\s+)?(?:row|record)\b"
    r")",
    re.IGNORECASE | re.DOTALL,
)


def _render_value(raw: str) -> str:
    """Convert a raw NL token into a SQL literal."""
    raw = raw.strip()
    if not raw:
        raise FQValidationError("Empty value in NL insert.")
    # User-supplied quotes/backticks → preserve as string with SQL-escaped inner.
    if len(raw) >= 2 and raw[0] == raw[-1] and raw[0] in ("'", '"', "`"):
        inner = raw[1:-1].replace("'", "''")
        return f"'{inner}'"
    low = raw.lower()
    if low == "null":
        return "NULL"
    if low == "true":
        return "TRUE"
    if low == "false":
        return "FALSE"
    if re.match(r'^-?\d+$', raw):
        return raw
    if re.match(r'^-?\d+\.\d+$', raw):
        return raw
    inner = raw.replace("'", "''")
    return f"'{inner}'"


def _nl_insert_to_sql(message: str) -> str:
    """Parse a business-user NL insert and emit a fully-qualified INSERT SQL.

    Raises FQValidationError on any missing/ambiguous piece (no guessing).
    """
    msg = message.strip().rstrip(".;!?")

    # 1. Resolve target — try inline dotted form first.
    catalog = schema = table = None
    m = re.search(
        r'(?:into|to)\s+([a-zA-Z_]\w*)\.([a-zA-Z_]\w*)\.([a-zA-Z_]\w*)',
        msg, re.IGNORECASE,
    )
    if m:
        catalog, schema, table = m.group(1), m.group(2), m.group(3)
    else:
        mc = re.search(r'\bcatalog\s+([a-zA-Z_]\w*)', msg, re.IGNORECASE)
        ms = re.search(r'\bschema\s+([a-zA-Z_]\w*)',  msg, re.IGNORECASE)
        mt = re.search(r'\btable\s+([a-zA-Z_]\w*)',   msg, re.IGNORECASE)
        catalog = mc.group(1) if mc else None
        schema  = ms.group(1) if ms else None
        table   = mt.group(1) if mt else None

    missing = [n for n, v in (("catalog", catalog), ("schema", schema), ("table", table)) if not v]
    if missing:
        raise FQValidationError(
            "Could not resolve " + ", ".join(missing) + " for natural-language insert. "
            "Specify all three explicitly, e.g. "
            "'add a new row into catalog mcp2ohio, schema test_writes, table demo "
            "with id 1, name foo, amount 9.99' "
            "— or use the dotted form 'into mcp2ohio.test_writes.demo'."
        )

    for part in (catalog, schema, table):
        try:
            StarburstClient.validate_identifier(part)
        except ValueError as e:
            raise FQValidationError(str(e))

    # 2. Resolve column-value pairs — segment after with|where|having.
    body = re.split(r'\b(?:with|where|having)\b', msg, maxsplit=1, flags=re.IGNORECASE)
    if len(body) < 2:
        raise FQValidationError(
            "Could not find column values. Add a 'with ...' clause listing column-value "
            "pairs, e.g. 'with id 1, name foo, amount 9.99'."
        )
    pairs_text = body[1].strip().rstrip(".;,!?")
    # Normalize "and" before the last comma-segment.
    pairs_text = re.sub(r',?\s*\band\b\s+', ', ', pairs_text, flags=re.IGNORECASE)

    raw_pairs = [p.strip() for p in pairs_text.split(',') if p.strip()]
    if not raw_pairs:
        raise FQValidationError("No column-value pairs found after 'with' / 'where'.")

    pair_re = re.compile(r'^([a-zA-Z_]\w*)\s*(?:=|:|is)?\s*(.+?)$')
    cols: list[str] = []
    vals: list[str] = []
    for raw in raw_pairs:
        m = pair_re.match(raw)
        if not m:
            raise FQValidationError(f"Could not parse column-value pair: '{raw}'")
        col = m.group(1)
        try:
            StarburstClient.validate_identifier(col)
        except ValueError as e:
            raise FQValidationError(str(e))
        cols.append(col)
        vals.append(_render_value(m.group(2)))

    if len(cols) != len(set(cols)):
        raise FQValidationError(f"Duplicate columns in NL insert: {cols}")

    fqn = f"{catalog}.{schema}.{table}"
    return f"INSERT INTO {fqn} ({', '.join(cols)}) VALUES ({', '.join(vals)})"


# ---------------------------------------------------------------------------
# Business-user NL UPDATE
#
# Recognizes phrasings like:
#   "In catalog X, schema Y, table Z, update the row where <cond> and set <col> to <val>"
#   "Change <col> to <val> in catalog X, schema Y, table Z for the row where <cond>"
#   "Update X.Y.Z and set <col> = <val> where <cond>"
# Requires explicit catalog + schema + table (or dotted form) AND a WHERE clause —
# no guessing on either object resolution or filter scope.
# ---------------------------------------------------------------------------

_NL_UPDATE_TRIGGER = re.compile(
    r"^\s*(?:"
    r"change\s+[a-zA-Z_]\w*\s+to\b|"
    r"update\s+(?:the\s+(?:row|record)|[a-zA-Z_]\w*(?:\.[a-zA-Z_]\w*){0,2}\s+and\s+set)|"
    r"in\s+catalog\s+\w+.*?\b(?:update|change)\b"
    r")",
    re.IGNORECASE | re.DOTALL,
)

_SET_PAIR_RE = re.compile(
    r'\b(?:and\s+)?set\s+([a-zA-Z_]\w*)\s*(?:=|\bto\b)\s*(.+?)'
    r'(?=\s+(?:where|for|and\s+(?:set|change))\b|,|\s*$)',
    re.IGNORECASE,
)
_CHANGE_PAIR_RE = re.compile(
    r'\bchange\s+([a-zA-Z_]\w*)\s+to\s+(.+?)'
    r'(?=\s+(?:in|for|where|and\s+(?:set|change))\b|,|\s*$)',
    re.IGNORECASE,
)
_WHERE_RE = re.compile(
    r'\bwhere\s+(.+?)(?=\s+and\s+(?:set|change)\b|\s*$)',
    re.IGNORECASE,
)


def _nl_update_to_sql(message: str) -> str:
    """Parse a business-user NL update and emit a fully-qualified UPDATE SQL.

    Raises FQValidationError on any missing/ambiguous piece (no guessing).
    """
    msg = message.strip().rstrip(".;!?")

    # 1. Resolve target — inline dotted form first.
    catalog = schema = table = None
    m = re.search(
        r'(?:update|on|to|into)\s+([a-zA-Z_]\w*)\.([a-zA-Z_]\w*)\.([a-zA-Z_]\w*)',
        msg, re.IGNORECASE,
    )
    if m:
        catalog, schema, table = m.group(1), m.group(2), m.group(3)
    else:
        mc = re.search(r'\bcatalog\s+([a-zA-Z_]\w*)', msg, re.IGNORECASE)
        ms = re.search(r'\bschema\s+([a-zA-Z_]\w*)',  msg, re.IGNORECASE)
        mt = re.search(r'\btable\s+([a-zA-Z_]\w*)',   msg, re.IGNORECASE)
        catalog = mc.group(1) if mc else None
        schema  = ms.group(1) if ms else None
        table   = mt.group(1) if mt else None

    missing = [n for n, v in (("catalog", catalog), ("schema", schema), ("table", table)) if not v]
    if missing:
        raise FQValidationError(
            "Could not resolve " + ", ".join(missing) + " for natural-language update. "
            "Specify all three explicitly, e.g. "
            "'in catalog mcp2ohio, schema test_writes, table demo, update the row where id=1 and set name to foo' "
            "— or use the dotted form 'update mcp2ohio.test_writes.demo and set ...'."
        )

    for part in (catalog, schema, table):
        try:
            StarburstClient.validate_identifier(part)
        except ValueError as e:
            raise FQValidationError(str(e))

    # 2. Resolve SET pairs — accept both 'set <col> to|= <val>' and 'change <col> to <val>'.
    set_pairs: list[tuple[str, str]] = []
    for sm in _CHANGE_PAIR_RE.finditer(msg):
        set_pairs.append((sm.group(1), sm.group(2).strip()))
    for sm in _SET_PAIR_RE.finditer(msg):
        set_pairs.append((sm.group(1), sm.group(2).strip()))

    if not set_pairs:
        raise FQValidationError(
            "UPDATE requires a SET clause specifying which columns to change. "
            "Add 'set <col> to <val>' or 'change <col> to <val>'."
        )

    # 3. Resolve WHERE clause — mandatory.
    wm = _WHERE_RE.search(msg)
    if not wm:
        wm = re.search(
            r'\bfor\s+the\s+(?:row|record)\s+where\s+(.+?)(?=\s+and\s+(?:set|change)\b|\s*$)',
            msg, re.IGNORECASE,
        )
    if not wm:
        raise FQValidationError(
            "UPDATE requires a WHERE clause to limit affected rows. "
            "Add 'where <condition>' (e.g. 'where id = 1')."
        )
    where_clause = wm.group(1).strip()
    if not where_clause:
        raise FQValidationError("UPDATE WHERE clause is empty.")

    # 4. Render columns + values; reject duplicates.
    cols_seen: set[str] = set()
    set_clauses: list[str] = []
    for col, raw_val in set_pairs:
        try:
            StarburstClient.validate_identifier(col)
        except ValueError as e:
            raise FQValidationError(str(e))
        if col in cols_seen:
            raise FQValidationError(f"Duplicate column in SET clause: '{col}'")
        cols_seen.add(col)
        set_clauses.append(f"{col} = {_render_value(raw_val)}")

    fqn = f"{catalog}.{schema}.{table}"
    return f"UPDATE {fqn} SET {', '.join(set_clauses)} WHERE {where_clause}"


def _nl_to_sql(message: str, catalog: str, schema: str) -> str | None:
    """Return SQL string or None if no pattern matched and not raw SQL."""
    msg = message.strip()
    low = msg.lower()

    # Business-user NL insert — runs first so it preempts the dot-passthrough
    # for inputs like "Insert a record into mcp2ohio.test_writes.demo where ...".
    if _NL_INSERT_TRIGGER.match(low):
        return _nl_insert_to_sql(msg)

    # Business-user NL update — same reason.
    if _NL_UPDATE_TRIGGER.match(low):
        return _nl_update_to_sql(msg)

    # Extract catalog/schema from NL context hints
    catalog, schema = _extract_context(low, catalog, schema)
    fq = f"{catalog}.{schema}"

    # Clean context hints from message for pattern matching
    clean = _clean_nl(low)

    # Exact NL matches first (before raw SQL check)
    if clean in ("show tables", "list tables", "tables", "show all tables", "list all tables"):
        return f"SHOW TABLES FROM {fq}"

    if clean in ("show schemas", "list schemas", "schemas", "show all schemas"):
        return f"SHOW SCHEMAS FROM {catalog}"

    # ---- DDL / DML NL patterns (must run before raw-SQL passthrough so that
    # numeric literals like "1.0" don't trigger the dot-passthrough on
    # writes; bare table names are auto-qualified to {fq}.{table}). --------

    def _qualify(tbl: str) -> str:
        return tbl if "." in tbl else f"{fq}.{tbl}"

    def _create_table_sql(table_name: str, columns: list[tuple[str, str]]) -> str:
        cols = ", ".join(f"{name} {dtype}" for name, dtype in columns)
        return f"CREATE TABLE IF NOT EXISTS {catalog}.{schema}.{table_name} ({cols})"

    # Governed business-user CREATE phrases. Keep this intentionally
    # deterministic: no arbitrary resource names or unsupported types.
    m = re.match(r"create\s+(?:a\s+)?(?:scratch\s+)?schema$", clean, re.IGNORECASE)
    if m:
        return f"CREATE SCHEMA IF NOT EXISTS {catalog}.scratch_sch"

    m = re.match(r"create\s+(?:a\s+)?sandbox\s+schema(?:\s+for\s+testing)?$", clean, re.IGNORECASE)
    if m:
        return f"CREATE SCHEMA IF NOT EXISTS {catalog}.sandbox_schema"

    m = re.match(r"create\s+(?:a\s+)?(?:demo\s+analytics\s+namespace|demo\s+analytics\s+schema)$", clean, re.IGNORECASE)
    if m:
        return f"CREATE SCHEMA IF NOT EXISTS {catalog}.demo_analytics"

    m = re.match(r"create\s+(?:a\s+)?staging\s+schema(?:\s+for\s+imports)?$", clean, re.IGNORECASE)
    if m:
        return f"CREATE SCHEMA IF NOT EXISTS {catalog}.staging_area"

    m = re.match(r"create\s+(?:a\s+)?(?:scratch|sandbox)\s+table(?:\s+for\s+dashboard\s+testing)?$", clean, re.IGNORECASE)
    if m:
        return _create_table_sql("scratch_nl_table", [("id", "INTEGER"), ("label", "VARCHAR")])

    m = re.match(r"create\s+(?:a\s+)?customer\s+(?:metrics|sales)\s+table(?:\s+with\s+id\s+and\s+revenue)?$", clean, re.IGNORECASE)
    if m:
        return _create_table_sql("customer_metrics_nl", [("customer_id", "BIGINT"), ("customer_name", "VARCHAR"), ("revenue", "DOUBLE"), ("growth_percent", "DOUBLE")])

    m = re.match(r"create\s+(?:a\s+)?(?:dashboard\s+metrics|chart\s+cache|dashboard\s+summary)\s+table$", clean, re.IGNORECASE)
    if m:
        return _create_table_sql("dashboard_metrics_nl", [("dashboard_id", "VARCHAR"), ("chart_type", "VARCHAR"), ("metric_value", "DOUBLE"), ("generated_at", "TIMESTAMP")])

    m = re.match(r"create\s+(?:an?\s+)?(?:executive\s+kpi|revenue\s+tracking)\s+table$", clean, re.IGNORECASE)
    if m:
        return _create_table_sql("executive_kpis_nl", [("kpi_name", "VARCHAR"), ("current_value", "DOUBLE"), ("target_value", "DOUBLE"), ("trend_direction", "VARCHAR")])

    m = re.match(r"create\s+(?:a\s+)?product\s+analytics\s+table$", clean, re.IGNORECASE)
    if m:
        return _create_table_sql("product_analytics_nl", [("product_id", "BIGINT"), ("product_name", "VARCHAR"), ("revenue", "DOUBLE"), ("rating", "DOUBLE")])

    m = re.match(r"delete\s+(?:rows?\s+)?(?:from\s+)?([\w.]+)(?:\s+where\s+(.+))?$", clean, re.IGNORECASE)
    if m:
        tbl, where = m.group(1), m.group(2)
        return f"DELETE FROM {_qualify(tbl)}" + (f" WHERE {where}" if where else "")

    m = re.match(r"truncate\s+(?:table\s+)?([\w.]+)$", clean, re.IGNORECASE)
    if m:
        return f"TRUNCATE TABLE {_qualify(m.group(1))}"

    m = re.match(r"drop\s+table\s+([\w.]+)$", clean, re.IGNORECASE)
    if m:
        return f"DROP TABLE {_qualify(m.group(1))}"

    m = re.match(r"drop\s+schema\s+([\w.]+)$", clean, re.IGNORECASE)
    if m:
        sch = m.group(1)
        return f"DROP SCHEMA {sch if '.' in sch else f'{catalog}.{sch}'}"

    m = re.match(r"create\s+schema\s+([\w.]+)$", clean, re.IGNORECASE)
    if m:
        sch = m.group(1)
        return f"CREATE SCHEMA {sch if '.' in sch else f'{catalog}.{sch}'}"

    m = re.match(r"update\s+([\w.]+)\s+set\s+(.+?)(?:\s+where\s+(.+))?$", clean, re.IGNORECASE)
    if m:
        tbl, sets, where = m.group(1), m.group(2), m.group(3)
        return f"UPDATE {_qualify(tbl)} SET {sets}" + (f" WHERE {where}" if where else "")

    m = re.match(r"insert\s+(?:into\s+)?([\w.]+)\s+(.+)$", clean, re.IGNORECASE)
    if m:
        tbl, rest = m.group(1), m.group(2)
        return f"INSERT INTO {_qualify(tbl)} {rest}"

    # Raw SQL passthrough — if it contains dots (qualified names), treat as SQL
    if _SQL_START.match(msg) and '.' in msg:
        return msg

    # describe / show columns
    m = re.match(r"(?:describe|show columns of|columns of)\s+(\w+)", clean)
    if m:
        return f"DESCRIBE {fq}.{m.group(1)}"

    # show all data / select from
    m = re.match(r"(?:show all data from|show data from|select from|show me data from|show me)\s+(\w+)", clean)
    if m:
        return f"SELECT * FROM {fq}.{m.group(1)} LIMIT 100"

    # count rows
    m = re.match(r"count (?:rows in|rows from|from|of)\s+(\w+)", clean)
    if m:
        return f"SELECT COUNT(*) FROM {fq}.{m.group(1)}"

    # top N column from table
    m = re.match(r"top\s+(\d+)\s+(\w+)\s+from\s+(\w+)", clean)
    if m:
        n, col, tbl = m.group(1), m.group(2), m.group(3)
        return f"SELECT * FROM {fq}.{tbl} ORDER BY {col} DESC LIMIT {n}"

    # aggregation: sum/avg/min/max of column from table
    m = re.match(r"(sum|avg|min|max|average)\s+(?:of\s+)?(\w+)\s+from\s+(\w+)", clean)
    if m:
        agg = m.group(1).upper()
        if agg == "AVERAGE":
            agg = "AVG"
        col, tbl = m.group(2), m.group(3)
        return f"SELECT {agg}({col}) FROM {fq}.{tbl}"

    # group by column from table
    m = re.match(r"group\s+by\s+(\w+)\s+(?:from|in)\s+(\w+)", clean)
    if m:
        col, tbl = m.group(1), m.group(2)
        return f"SELECT {col}, COUNT(*) FROM {fq}.{tbl} GROUP BY {col}"

    # Fallback: raw SQL without dots
    if _SQL_START.match(msg):
        return msg

    return None


_nl2sql_service = NL2SQLService(provider=_llm_provider, legacy_translator=_nl_to_sql)


# ---------------------------------------------------------------------------
def _try_other_schemas(message: str, failed_sql: str, catalog: str, default_schema: str):
    """When a table isn't found in the default schema, search cached schema data."""
    m = re.search(r'FROM\s+\S+\.(\w+)\.(\w+)', failed_sql, re.IGNORECASE)
    if not m:
        return None
    table_name = m.group(2).lower()

    # Use cached schema to avoid slow Starburst round-trips
    cached = _schema_cache.get("data")
    if cached and cached.get("schemas"):
        for s in cached["schemas"]:
            if s["name"] == default_schema or s["name"] == "information_schema":
                continue
            for t in s.get("tables", []):
                if t["name"].lower() == table_name:
                    return failed_sql.replace(
                        f"{catalog}.{default_schema}.{m.group(2)}",
                        f"{catalog}.{s['name']}.{t['name']}"
                    )

    # Fallback: single direct query if cache is empty
    try:
        result = client.execute(
            f"SELECT table_schema, table_name FROM {catalog}.information_schema.tables "
            f"WHERE table_name = '{table_name}' AND table_schema != '{default_schema}' LIMIT 1"
        )
        if result.get("rows"):
            found_schema = result["rows"][0][0]
            found_table = result["rows"][0][1]
            return failed_sql.replace(
                f"{catalog}.{default_schema}.{m.group(2)}",
                f"{catalog}.{found_schema}.{found_table}"
            )
    except Exception:
        pass
    return None


# ---------------------------------------------------------------------------
# Endpoints
# ---------------------------------------------------------------------------

@app.get("/")
async def index():
    index_file = HERE / "index.html"
    if not index_file.exists():
        raise HTTPException(status_code=404, detail="index.html not found")
    return FileResponse(str(index_file))


_schema_cache: dict = {"data": None, "ts": 0}
_SCHEMA_TTL = 300  # 5 minutes

def _describe_table(catalog: str, schema_name: str, table_name: str) -> dict:
    """Fetch columns for one table. Runs in a thread."""
    columns = []
    try:
        col_result = client.execute(f"DESCRIBE {catalog}.{schema_name}.{table_name}")
        for crow in col_result.get("rows", []):
            columns.append({"name": crow[0], "type": crow[1] if len(crow) > 1 else "unknown"})
    except Exception:
        pass
    return {"name": table_name, "columns": columns}


def _fetch_schema():
    catalog = client.catalog or "mcp2ohio"

    # Step 1: Get all schemas
    schema_result = client.execute(f"SHOW SCHEMAS FROM {catalog}")
    schema_names = [row[0] for row in schema_result.get("rows", [])]

    # Step 2: Get all tables for all schemas in parallel
    schema_tables = {}
    with ThreadPoolExecutor(max_workers=len(schema_names)) as pool:
        def get_tables(s):
            try:
                r = client.execute(f"SHOW TABLES FROM {catalog}.{s}")
                return s, [row[0] for row in r.get("rows", [])]
            except Exception:
                return s, []
        futures = [pool.submit(get_tables, s) for s in schema_names]
        for f in as_completed(futures):
            name, tables = f.result()
            schema_tables[name] = tables

    # Step 3: Describe ALL tables across ALL schemas in parallel
    all_tasks = []
    for schema_name, table_names in schema_tables.items():
        for table_name in table_names:
            all_tasks.append((schema_name, table_name))

    table_data = {}  # (schema, table) -> {name, columns}
    with ThreadPoolExecutor(max_workers=min(len(all_tasks), 10)) as pool:
        futures = {
            pool.submit(_describe_table, catalog, s, t): (s, t)
            for s, t in all_tasks
        }
        for f in as_completed(futures):
            s, t = futures[f]
            table_data[(s, t)] = f.result()

    # Step 4: Assemble result
    schemas = []
    for schema_name in sorted(schema_tables.keys()):
        tables = [table_data[(schema_name, t)] for t in schema_tables[schema_name]]
        schemas.append({"name": schema_name, "tables": tables})
    return {"catalog": catalog, "schemas": schemas}

@app.get("/api/schema")
async def get_schema(refresh: int = 0):
    try:
        now = time.time()
        if not refresh and _schema_cache["data"] and (now - _schema_cache["ts"]) < _SCHEMA_TTL:
            return _schema_cache["data"]
        result = _fetch_schema()
        _schema_cache["data"] = result
        _schema_cache["ts"] = now
        return result
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc))


@app.post("/api/query")
async def run_query(req: QueryRequest):
    try:
        return _exec(req.sql)
    except QueryError as exc:
        raise HTTPException(status_code=400, detail=str(exc))


@app.post("/api/chat")
async def chat(req: ChatRequest):
    catalog = _catalog(req.context)
    schema = _schema(req.context)

    try:
        nl_result = _nl2sql_service.translate(
            req.message,
            catalog,
            schema,
            _schema_cache.get("data"),
        )
        sql = nl_result.sql
    except FQValidationError as exc:
        raise HTTPException(
            status_code=400,
            detail={"error": "fq_validation", "message": str(exc), "sql": None},
        )
    if sql is None:
        if nl_result.source == "rejected":
            return {
                "sql": None,
                "columns": [],
                "rows": [],
                "row_count": 0,
                "chart_suggestion": "table",
                "message": "NO RELEVANT INFORMATION EXISTS IN THE ENTERPRISE DATA SOURCES.",
                "llm_source": nl_result.source,
                "query_type": nl_result.query_type,
                "query_origin": "enterprise_grounding_policy",
                "data_grounded": False,
                "source_type": "approved_enterprise_sources_only",
                "data_sources": [],
                "grounding_reason": nl_result.grounding_reason,
                "warnings": [
                    "The request was not answered because it did not match approved enterprise data assets."
                ],
            }
        if nl_result.source == "llm_error":
            detail = nl_result.error or "The configured LLM provider did not return usable SQL."
            return {
                "sql": None,
                "columns": [],
                "rows": [],
                "row_count": 0,
                "chart_suggestion": "table",
                "message": (
                    "The LLM provider is configured, but it could not generate executable SQL. "
                    f"Provider detail: {detail}"
                ),
                "llm_source": nl_result.source,
                "query_type": nl_result.query_type,
                "query_origin": "enterprise_grounded_nl2sql",
                "data_grounded": False,
                "source_type": "approved_enterprise_sources_only",
                "data_sources": nl_result.matched_terms,
                "grounding_reason": nl_result.grounding_reason,
                "warnings": nl_result.warnings,
            }
        return {
            "sql": None,
            "columns": [],
            "rows": [],
            "row_count": 0,
            "chart_suggestion": "table",
            "message": (
                "I couldn't understand that query. Try phrases like:\n"
                "- show tables\n"
                "- describe <table>\n"
                "- show all data from <table>\n"
                "- count rows in <table>\n"
                "- top 10 <column> from <table>\n"
                "- sum/avg/min/max of <column> from <table>\n"
                "- group by <column> from <table>\n"
                "- Or enter raw SQL directly."
            ),
            "llm_source": "none",
            "query_type": "unmatched",
            "data_grounded": False,
            "source_type": "approved_enterprise_sources_only",
            "data_sources": [],
        }

    # Classify, validate FQ for DDL/DML, check permissions, gate destructive ops.
    op_label, perm_key = _classify_sql(sql)
    confirm = bool((req.context or {}).get("confirm", False))
    try:
        target = _validate_fq(sql, op_label)
        _check_permission(perm_key)
        _enforce_destructive_confirm(perm_key, op_label, target, confirm, sql)
    except FQValidationError as exc:
        raise HTTPException(status_code=400, detail={"error": "fq_validation", "message": str(exc), "sql": sql})
    except PermissionDenied as exc:
        raise HTTPException(status_code=403, detail={"error": "permission_denied", "message": str(exc), "sql": sql})
    except ConfirmRequired as exc:
        return {
            "sql": sql,
            "columns": [],
            "rows": [],
            "row_count": 0,
            "chart_suggestion": "table",
            "requires_confirm": True,
            "operation": op_label,
            "target": target,
            "message": str(exc),
        }

    try:
        result = _exec(sql)
        result = _attach_write_verification(result, sql, op_label, target)
    except QueryError as exc:
        err_str = str(exc)
        if perm_key == "read" and ("TABLE_NOT_FOUND" in err_str or "does not exist" in err_str):
            alt_sql = _try_other_schemas(req.message, sql, catalog, schema)
            if alt_sql:
                sql = alt_sql
                result = _exec(sql)
            else:
                raise HTTPException(status_code=400, detail=err_str)
        else:
            raise HTTPException(status_code=400, detail=err_str)

    suggested = _suggest_chart(result["columns"], result["rows"])
    chart = _final_chart_suggestion(nl_result.chart, suggested)
    insights = build_result_insights(
        message=req.message,
        sql=sql,
        columns=result["columns"],
        rows=result["rows"],
        chart=chart,
    )
    payload = {
        "sql": sql,
        "columns": result["columns"],
        "rows": result["rows"],
        "row_count": result["row_count"],
        "chart_suggestion": chart,
        "operation": op_label,
        "message": f"Executed: {sql}",
        "summary": insights.summary,
        "insights": insights.insights,
        "dashboard": insights.dashboard,
        "kpis": insights.kpis,
        "trend_highlights": insights.trend_highlights,
        "recommendations": insights.recommendations,
        "chart_explanation": insights.chart_explanation,
        "llm_source": nl_result.source,
        "query_type": nl_result.query_type,
        "query_origin": "direct_sql" if nl_result.source == "sql" else "enterprise_grounded_nl2sql",
        "data_grounded": True,
        "source_type": "approved_enterprise_sources_only",
        "data_sources": _data_sources_for_payload(sql, nl_result.matched_terms),
        "grounding_reason": nl_result.grounding_reason,
        "confidence": nl_result.confidence,
        "assumptions": nl_result.assumptions,
        "followups": nl_result.followups,
        "warnings": nl_result.warnings,
    }
    target_name = f"{target.get('catalog')}.{target.get('schema')}" + (f".{target['table']}" if target.get('table') else "")
    if result.get("verification_sql"):
        payload["verification_sql"] = result["verification_sql"]
        payload["verification_row_count"] = result.get("verification_row_count", result["row_count"])
    if result.get("verification_reason"):
        payload["verification_reason"] = result["verification_reason"]
    if "rows_affected" in result:
        payload["rows_affected"] = result["rows_affected"]
        payload["status"] = result.get("status", "ok")
        payload["message"] = f"{op_label} completed on {target_name} — rows_affected={result['rows_affected']}"
    elif op_label in {"CREATE TABLE", "CREATE SCHEMA"}:
        payload["status"] = result.get("status", "ok")
        payload["message"] = f"{op_label} completed on {target_name}"
    if result.get("verification_sql"):
        payload["message"] += f". Verification query returned {payload['row_count']} row(s)."
    if result.get("verification_reason"):
        payload["message"] += f". Verification query unavailable: {result['verification_reason']}"
    return payload


def _data_sources_for_payload(sql: str, matched_terms: list[str]) -> list[str]:
    sources = set(matched_terms or [])
    for cat, sch, tbl in re.findall(r'\b(?:FROM|JOIN|DESCRIBE)\s+([a-zA-Z_]\w*)\.([a-zA-Z_]\w*)\.([a-zA-Z_]\w*)', sql, re.IGNORECASE):
        sources.add(f"{cat}.{sch}.{tbl}")
    return sorted(sources)[:10]


@app.post("/api/export/{fmt}")
async def export_data(fmt: str, req: ExportRequest):
    if fmt not in ("csv", "xlsx", "html"):
        raise HTTPException(status_code=400, detail=f"Unsupported format: {fmt}")

    safe_title = re.sub(r"[^\w\-]", "_", req.title)

    if fmt == "csv":
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(req.columns)
        writer.writerows(req.rows)
        return Response(
            content=buf.getvalue(),
            media_type="text/csv",
            headers={"Content-Disposition": f'attachment; filename="{safe_title}.csv"'},
        )

    if fmt == "xlsx":
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = req.title[:31]
        ws.append(req.columns)
        for row in req.rows:
            ws.append(row)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{safe_title}.xlsx"'},
        )

    # html
    rows_html = "".join(
        "<tr>" + "".join(f"<td>{v}</td>" for v in row) + "</tr>" for row in req.rows
    )
    header_html = "".join(f"<th>{c}</th>" for c in req.columns)
    html = (
        "<!DOCTYPE html><html><head><meta charset='utf-8'>"
        f"<title>{req.title}</title>"
        "<style>table{border-collapse:collapse;width:100%}"
        "th,td{border:1px solid #ccc;padding:8px;text-align:left}"
        "th{background:#f5f5f5}</style></head><body>"
        f"<h2>{req.title}</h2>"
        f"<table><thead><tr>{header_html}</tr></thead>"
        f"<tbody>{rows_html}</tbody></table></body></html>"
    )
    return Response(
        content=html,
        media_type="text/html",
        headers={"Content-Disposition": f'attachment; filename="{safe_title}.html"'},
    )


# ---------------------------------------------------------------------------
# Startup
# ---------------------------------------------------------------------------

@app.on_event("startup")
async def startup_banner():
    print("=" * 60)
    print("  Starburst Galaxy Chatbot")
    print(f"  Host:    {client.host}")
    print(f"  Catalog: {client.catalog}")
    print(f"  Schema:  {client.schema}")
    print(f"  Auth:    {client.auth_mode}")
    print("=" * 60)
    # Pre-warm schema cache in background thread (non-blocking)
    def _warm():
        try:
            _schema_cache["data"] = _fetch_schema()
            _schema_cache["ts"] = time.time()
            print("  Schema cache warmed ✓")
        except Exception as e:
            print(f"  Schema cache warm failed: {e}")
    threading.Thread(target=_warm, daemon=True).start()


if __name__ == "__main__":
    import uvicorn
    uvicorn.run("app_jwt:app", host="0.0.0.0", port=8000, reload=True)
